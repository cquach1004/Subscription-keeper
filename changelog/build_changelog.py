"""Build change-log.xlsx from changes.csv.

changes.csv is the source of truth: add one row per change in each new version,
then run:  python3 build_changelog.py   (from this folder or with paths as arguments)
"""
import csv, os, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

here = os.path.dirname(os.path.abspath(__file__))
src = sys.argv[1] if len(sys.argv) > 1 else os.path.join(here, "changes.csv")
out = sys.argv[2] if len(sys.argv) > 2 else os.path.join(here, "change-log.xlsx")

with open(src, newline="") as f:
    rows = list(csv.reader(f))
head, body = rows[0], [r for r in rows[1:] if any(c.strip() for c in r)]

FONT = "Arial"
navy = PatternFill("solid", fgColor="1F3A5F")
band = PatternFill("solid", fgColor="F2F5F9")
line = Border(bottom=Side(style="thin", color="D5DCE6"))
widths = {"Version": 11, "Date (UTC)": 12, "Change": 70, "Requested by": 22, "Area / Page": 24,
          "Type": 10, "On public site": 20, "Status": 16, "Notes": 45}

wb = Workbook()
ws = wb.active
ws.title = "Change Log"
ws.append(head)
for c in ws[1]:
    c.font = Font(name=FONT, bold=True, color="FFFFFF")
    c.fill = navy
    c.alignment = Alignment(vertical="center", wrap_text=True)
ws.row_dimensions[1].height = 22
for i, r in enumerate(body, start=2):
    ws.append(r)
    for c in ws[i]:
        c.font = Font(name=FONT, size=10)
        c.alignment = Alignment(vertical="top", wrap_text=True)
        c.border = line
        if i % 2 == 0:
            c.fill = band
for i, h in enumerate(head, start=1):
    ws.column_dimensions[ws.cell(1, i).column_letter].width = widths.get(h, 16)
ws.freeze_panes = "B2"
ws.auto_filter.ref = ws.dimensions

info = wb.create_sheet("About")
lines = [
    ("Subscription Keeper change log", True),
    ("One row per change, tied to the app version it shipped in. Newest versions are at the bottom.", False),
    ("", False),
    ("Columns", True),
    ("Version: the number shown in the app's profile menu. '2–20' and '(est.)' mark early versions shipped before the version label existed (v39), so the exact number wasn't recorded.", False),
    ("Requested by: who asked for it. '(approved)' = Cuong approved a researched change; '(tester feedback)' = Cuong passing on testers' comments.", False),
    ("On public site: whether the version went to https://cquach1004.github.io/Subscription-keeper/ (the public site started at v41).", False),
    ("Status: Current, Changed later, Hidden or Removed in a later version.", False),
    ("", False),
    ("Adding new versions", True),
    ("Don't edit this spreadsheet by hand: it is rebuilt from changes.csv. Add a row to changes.csv for each change in the new version, then run build_changelog.py.", False),
    ("Master copy: the project's shared folder changelog/; a copy is kept in the GitHub repo under changelog/.", False),
]
for text, bold in lines:
    info.append([text])
    info.cell(info.max_row, 1).font = Font(name=FONT, bold=bold, size=12 if bold else 10)
    info.cell(info.max_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
info.column_dimensions["A"].width = 110

wb.save(out)
print(f"{len(body)} rows -> {out}")
